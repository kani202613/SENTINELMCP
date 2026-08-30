"""
SentinelMCP Full 127-Trace Benchmark Evaluation Harness (Milestone 5)
Evaluates SentinelMCP against 2 baseline security architectures:
1. Static MCP (Unprotected baseline)
2. MCP-Secure Lite (Static RBAC / pattern matching baseline)

Replays 127 trace sessions across:
- Traces - Main (81 sessions)
- Traces - Set B (Red-Team) (28 sessions)
- Traces - Evasion (18 sessions)

Computes Precision, Recall, F1-Score, Accuracy, FPR, and Per-Category Recall.
Writes empirical results directly into 'Results Template' sheet in SentinelMCP_Dataset_Clean.xlsx.
"""
import sys
import os
import json
import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from sentinel.interceptor import SentinelInterceptor
from baseline.static_mcp import run_static_mcp
from baseline.mcp_secure_lite import run_mcp_secure_lite

def normalize_tool_request(step: dict) -> tuple:
    raw_tool = str(step.get("tool_name", "file_tool")).lower()
    args = step.get("args", {})
    action = str(step.get("action", args.get("action", ""))).lower()

    if raw_tool in ["file_reader", "file_writer", "file_client"]:
        tool_name = "file_tool"
        action = action or ("read_file" if "read" in raw_tool else "write_file")
    elif raw_tool in ["email_client", "http_client", "api_client", "http_tool"]:
        tool_name = "http_tool"
        action = action or "send_request"
    elif raw_tool in ["github_client", "github_tool"]:
        tool_name = "github_tool"
        action = action or "read_issue"
    elif raw_tool in ["web_browser", "web_tool"]:
        tool_name = "web_tool"
        action = action or "fetch_page"
    elif raw_tool in ["database_client", "sql_client", "database_tool"]:
        tool_name = "database_tool"
        action = action or "query_select"
    elif raw_tool in ["pdf_reader", "pdf_client"]:
        tool_name = "pdf_reader"
        action = action or "read_pdf"
    elif raw_tool in ["slack_client", "slack_tool"]:
        tool_name = "slack_tool"
        action = action or "send_message"
    else:
        tool_name = "file_tool"
        action = action or "read_file"

    return tool_name, action

def evaluate_system_on_dataset(system_name: str, eval_func, trace_sessions: list) -> dict:
    tp, fp, tn, fn = 0, 0, 0, 0
    category_totals = {}
    category_hits = {}

    for trace in trace_sessions:
        session_id = trace["session_id"]
        is_attack = (trace["label"] == "ATTACK")
        category = trace["category"]

        if is_attack:
            category_totals[category] = category_totals.get(category, 0) + 1

        user_prompt = trace["user_prompt"]
        steps = trace["steps"]

        max_sri = 0
        final_decision = "SAFE"

        interceptor = SentinelInterceptor(audit_log_path="data/audit_log.json") if system_name == "SentinelMCP" else None

        for step in steps:
            tool_name, action = normalize_tool_request(step)
            args = step.get("args", {})
            source_trust = str(step.get("source_trust", "INTERNAL"))

            req = {
                "tool_name": tool_name,
                "action": action,
                "user_role": "junior_analyst",
                "context": user_prompt,
                "args": args,
                "session_id": session_id,
                "source_trust": source_trust
            }

            if system_name == "Static MCP":
                res = run_static_mcp(req)
            elif system_name == "MCP-Secure Lite":
                res = run_mcp_secure_lite(req)
            else:
                res = interceptor.intercept_and_execute(req)

            dec = res["decision"]
            sri = res.get("sri") or res.get("sri_details", {}).get("sri", 0)

            if dec in ["SUSPICIOUS", "BLOCKED"] or sri > max_sri:
                max_sri = sri
                final_decision = dec

        system_flagged = final_decision in ["SUSPICIOUS", "BLOCKED"]

        if is_attack:
            if system_flagged:
                tp += 1
                category_hits[category] = category_hits.get(category, 0) + 1
            else:
                fn += 1
        else:
            if system_flagged:
                fp += 1
            else:
                tn += 1

    total = tp + fp + tn + fn
    accuracy = (tp + tn) / total if total > 0 else 0.0
    precision = tp / (tp + fp) if (tp + fp) > 0 else 0.0
    recall = tp / (tp + fn) if (tp + fn) > 0 else 0.0
    f1 = (2 * precision * recall) / (precision + recall) if (precision + recall) > 0 else 0.0
    fpr = fp / (fp + tn) if (fp + tn) > 0 else 0.0

    category_recalls = {}
    for cat, tot in category_totals.items():
        hits = category_hits.get(cat, 0)
        category_recalls[cat] = hits / tot if tot > 0 else 0.0

    return {
        "system_name": system_name,
        "accuracy": accuracy,
        "precision": precision,
        "recall": recall,
        "f1": f1,
        "fpr": fpr,
        "tp": tp,
        "fp": fp,
        "tn": tn,
        "fn": fn,
        "category_recalls": category_recalls
    }

def run_full_evaluation():
    excel_path = "SentinelMCP_Dataset_Clean.xlsx"
    xl = pd.ExcelFile(excel_path)

    all_traces = []
    for sheet in ["Traces - Main", "Traces - Set B (Red-Team)", "Traces - Evasion"]:
        if sheet in xl.sheet_names:
            df = pd.read_excel(excel_path, sheet_name=sheet)
            for idx, r in df.iterrows():
                try:
                    steps = json.loads(str(r['Raw Trace Details']))
                    all_traces.append({
                        "session_id": str(r['Session ID']),
                        "label": str(r['Label']).upper(),
                        "category": str(r.get('Attack Category', 'Benign Workflow')),
                        "user_prompt": str(r['User Prompt']),
                        "steps": steps
                    })
                except Exception:
                    pass

    print("==================================================================================")
    print("      SentinelMCP Full 127-Trace Replay Evaluation Harness Suite")
    print("==================================================================================\n")
    print(f"Total Trace Sessions Loaded: {len(all_traces)} across 3 benchmark sheets.\n")

    m_static = evaluate_system_on_dataset("Static MCP", run_static_mcp, all_traces)
    m_lite = evaluate_system_on_dataset("MCP-Secure Lite", run_mcp_secure_lite, all_traces)
    m_sentinel = evaluate_system_on_dataset("SentinelMCP", None, all_traces)

    print("==================================================================================")
    print("                 OVERALL EVALUATION RESULTS TABLE (127 TRACES)")
    print("==================================================================================")
    print(f"{'System':>16} {'Accuracy (%)':>13} {'Precision (%)':>14} {'Recall (%)':>11} {'F1-Score (%)':>13} {'FPR (%)':>8} {'TP':>3} {'FP':>3} {'TN':>3} {'FN':>3}")
    for m in [m_static, m_lite, m_sentinel]:
        print(f"{m['system_name']:>16} {m['accuracy']*100:>12.2f}% {m['precision']*100:>13.2f}% {m['recall']*100:>10.2f}% {m['f1']*100:>12.2f}% {m['fpr']*100:>7.2f}% {m['tp']:>3} {m['fp']:>3} {m['tn']:>3} {m['fn']:>3}")

    print("\n\n==================================================================================")
    print("               PER-CATEGORY ATTACK RECALL BREAKDOWN TABLE")
    print("==================================================================================")
    print(f"{'Attack Category':>32} {'Static MCP Recall':>18} {'MCP-Secure Lite Recall':>23} {'SentinelMCP Recall':>19}")

    cats = list(m_sentinel["category_recalls"].keys())
    for cat in cats:
        r_stat = m_static["category_recalls"].get(cat, 0.0) * 100
        r_lite = m_lite["category_recalls"].get(cat, 0.0) * 100
        r_sent = m_sentinel["category_recalls"].get(cat, 0.0) * 100
        print(f"{cat:>32} {r_stat:>17.1f}% {r_lite:>22.1f}% {r_sent:>18.1f}%")

    write_results_to_excel(excel_path, m_static, m_lite, m_sentinel)
    return [m_static, m_lite, m_sentinel]

def write_results_to_excel(excel_path: str, m_static: dict, m_lite: dict, m_sentinel: dict):
    wb = load_workbook(excel_path)
    if 'Results Template' in wb.sheetnames:
        ws = wb['Results Template']
        ws['A2'] = "Static MCP"
        ws['B2'] = round(m_static['accuracy'] * 100, 2)
        ws['C2'] = round(m_static['precision'] * 100, 2)
        ws['D2'] = round(m_static['recall'] * 100, 2)
        ws['E2'] = round(m_static['f1'] * 100, 2)
        ws['F2'] = round(m_static['fpr'] * 100, 2)

        ws['A3'] = "MCP-Secure Lite"
        ws['B3'] = round(m_lite['accuracy'] * 100, 2)
        ws['C3'] = round(m_lite['precision'] * 100, 2)
        ws['D3'] = round(m_lite['recall'] * 100, 2)
        ws['E3'] = round(m_lite['f1'] * 100, 2)
        ws['F3'] = round(m_lite['fpr'] * 100, 2)

        ws['A4'] = "SentinelMCP"
        ws['B4'] = round(m_sentinel['accuracy'] * 100, 2)
        ws['C4'] = round(m_sentinel['precision'] * 100, 2)
        ws['D4'] = round(m_sentinel['recall'] * 100, 2)
        ws['E4'] = round(m_sentinel['f1'] * 100, 2)
        ws['F4'] = round(m_sentinel['fpr'] * 100, 2)

        wb.save(excel_path)
        print("\n\nSuccessfully written empirical evaluation metrics into 'Results Template' sheet in SentinelMCP_Dataset_Clean.xlsx!\n")

if __name__ == "__main__":
    run_full_evaluation()
